"""Shared helpers for importers: read xlsx/csv tabs into dicts keyed by header."""

from __future__ import annotations

import csv
import re
from datetime import date
from pathlib import Path
from typing import Any, Iterable

from gcdc.identity import normalise_email, split_emails
from gcdc.timeutil import parse_date


def _norm_header(h: Any) -> str:
    return re.sub(r"\s+", " ", str(h or "")).strip()


def read_table(path: str | Path, sheet: str | None = None, *, header_contains: Iterable[str] = (),
               max_scan: int = 12) -> list[dict[str, Any]]:
    """Read one sheet (xlsx) or file (csv). The header row is the first of the
    top rows that contains every name in header_contains (or the first row
    with 3+ filled cells when none are given). Empty rows are dropped."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as fh:
            rows = [list(r) for r in csv.reader(fh)]
    else:
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover - dependency message
            raise SystemExit("Reading .xlsx needs openpyxl:  pip install -e .[import]") from exc
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        if sheet is None:
            ws = wb.worksheets[0]
        else:
            matches = [ws for ws in wb.worksheets if ws.title.strip().lower() == sheet.strip().lower()]
            if not matches:
                raise ValueError(f"{path.name}: no sheet named {sheet!r} (have: {', '.join(wb.sheetnames)})")
            ws = matches[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wanted = [w.lower() for w in header_contains]
    header_idx = None
    for i, row in enumerate(rows[:max_scan]):
        cells = [_norm_header(c).lower() for c in row]
        if wanted and all(any(w == c or c.startswith(w) for c in cells) for w in wanted):
            header_idx = i
            break
        if not wanted and sum(1 for c in cells if c) >= 3:
            header_idx = i
            break
    if header_idx is None:
        raise ValueError(f"{path.name}{' / ' + sheet if sheet else ''}: header row with {list(header_contains)} not found")
    header = [_norm_header(h) for h in rows[header_idx]]
    out = []
    for line_no, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
        if not any(v not in (None, "") for v in row):
            continue
        rec = {}
        for h, v in zip(header, row):
            if h and h not in rec:
                rec[h] = v.strip() if isinstance(v, str) else v
        rec["_row"] = line_no
        out.append(rec)
    return out


def sheet_names(path: str | Path) -> list[str]:
    import openpyxl
    return openpyxl.load_workbook(path, read_only=True).sheetnames


def col(rec: dict[str, Any], *names: str) -> Any:
    """First non-empty value among the named columns. An exact header match wins;
    a prefix match (e.g. 'Notes' for 'Notes — what they said') is used only when
    no header matches exactly, so 'Email' never falls through to 'Email 1 sent'."""
    for name in names:
        low = name.lower()
        exact = [k for k in rec if k.lower() == low]
        keys = exact or [k for k in rec if k.lower().startswith(low)]
        for key in keys:
            value = rec[key]
            if value not in (None, ""):
                return value
    return None


def yes(value: Any) -> int | None:
    if value in (None, ""):
        return None
    text = str(value).strip().lower()
    if text.startswith(("y", "true", "1", "verified")):
        return 1
    if text.startswith(("n", "false", "0")):
        return 0
    return None


def first_email(value: Any) -> str | None:
    emails = split_emails(value)
    return emails[0] if emails else normalise_email(value)


def safe_date(value: Any, *, default_year: int = 2026, not_before: date | None = None) -> tuple[date | None, str | None]:
    """Parse a sheet date. Returns (date, problem). Dates before not_before are rejected, not guessed."""
    if value in (None, ""):
        return None, None
    try:
        d = parse_date(value, default_year=default_year)
    except ValueError as exc:
        return None, str(exc)
    if d and not_before and d < not_before:
        return None, f"date {value!r} is before records start {not_before.isoformat()} — not imported"
    return d, None


def as_int(value: Any) -> int | None:
    try:
        return int(float(value)) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None
